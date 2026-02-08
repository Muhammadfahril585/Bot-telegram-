Import os
import io
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import (
    ContextTypes,
    ConversationHandler,
    CommandHandler,
    MessageHandler,
    CallbackQueryHandler,
    filters
)
from docx import Document

# Definisi State untuk Conversation
MENUNGGU_KATA_LAMA, KONFIRMASI_AKSI, MENUNGGU_KATA_BARU, PILIHAN_LANJUT = range(4)

async def mulai_edit_v2(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.message.reply_to_message or not update.message.reply_to_message.document:
        await update.message.reply_text("âŒ **Gagal:** Silakan **Balas (Reply)** file Word (.docx) lalu ketik `/edit`.")
        return ConversationHandler.END

    # Download file pertama kali dan simpan di memori (BytesIO)
    file_doc = await context.bot.get_file(update.message.reply_to_message.document.file_id)
    file_content = io.BytesIO()
    await file_doc.download_to_memory(file_content)

    # Simpan dokumen ke user_data agar bisa diedit berkali-kali
    context.user_data['doc_buffer'] = file_content
    context.user_data['file_name'] = update.message.reply_to_message.document.file_name

    await update.message.reply_text("ðŸ” Masukkan **Kata Lama** yang ingin Anda cari:")
    return MENUNGGU_KATA_LAMA

async def proses_cari_kata(update: Update, context: ContextTypes.DEFAULT_TYPE):
    kata_lama = update.message.text
    context.user_data['kata_lama'] = kata_lama

    # Baca dokumen dari buffer
    doc_buffer = context.user_data['doc_buffer']
    doc_buffer.seek(0)
    doc = Document(doc_buffer)

    # Hitung kemunculan kata
    count = sum(p.text.count(kata_lama) for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                count += sum(p.text.count(kata_lama) for p in cell.paragraphs)

    if count == 0:
        await update.message.reply_text(f"âš ï¸ Kata '{kata_lama}' tidak ditemukan. Silakan masukkan kata lain:")
        return MENUNGGU_KATA_LAMA

    keyboard = [
        [InlineKeyboardButton(f"Ganti Semua ({count})", callback_data="ganti_semua")],
        [InlineKeyboardButton("Batal", callback_data="batal")]
    ]

    await update.message.reply_text(
        f"âœ… Ditemukan **{count}** kata '{kata_lama}'.\nSilakan pilih opsi:",
        reply_markup=InlineKeyboardMarkup(keyboard)
    )
    return KONFIRMASI_AKSI

async def minta_kata_baru(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()

    if query.data == "batal":
        await query.edit_message_text("âŒ Edit dibatalkan.")
        return ConversationHandler.END

    await query.edit_message_text(f"Ok, masukkan **Kata Baru** untuk mengganti '{context.user_data['kata_lama']}':")
    return MENUNGGU_KATA_BARU

async def eksekusi_simpan_perubahan(update: Update, context: ContextTypes.DEFAULT_TYPE):
    kata_baru = update.message.text
    kata_lama = context.user_data['kata_lama']

    # Ambil dokumen dari buffer
    doc_buffer = context.user_data['doc_buffer']
    doc_buffer.seek(0)
    doc = Document(doc_buffer)

    # Proses Edit (Format Safe)
    def replace_text(paragraphs):
        for p in paragraphs:
            if kata_lama in p.text:
                for run in p.runs:
                    if kata_lama in run.text:
                        run.text = run.text.replace(kata_lama, kata_baru)

    replace_text(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                replace_text(cell.paragraphs)

    # Simpan kembali ke buffer
    new_buffer = io.BytesIO()
    doc.save(new_buffer)
    context.user_data['doc_buffer'] = new_buffer

    # Kirim tombol pilihan lanjut
    keyboard = [
        [InlineKeyboardButton("âž• Lanjut Edit Kata Lain", callback_data="lanjut_edit")],
        [InlineKeyboardButton("âœ… Cukup & Kirim File", callback_data="cukup_kirim")]
    ]

    await update.message.reply_text(
        f"ðŸ’¾ Perubahan disimpan! '{kata_lama}' âž¡ï¸ '{kata_baru}'.\n\nApa langkah selanjutnya?",
        reply_markup=InlineKeyboardMarkup(keyboard)
    )
    return PILIHAN_LANJUT
async def handle_pilihan_akhir(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()

    if query.data == "lanjut_edit":
        await query.edit_message_text("ðŸ” Silakan masukkan **Kata Lama** berikutnya:")
        return MENUNGGU_KATA_LAMA
    else:
        # 1. Ubah teks tombol menjadi status loading
        # Kita simpan ke variabel status agar bisa dikelola
        status_msg = await query.edit_message_text("â³ Sedang memproses dan mengirim file terbaru...")

        try:
            doc_buffer = context.user_data['doc_buffer']
            doc_buffer.seek(0)

            # 2. Kirim file final
            await query.message.reply_document(
                document=doc_buffer,
                filename=f"Revisi_{context.user_data.get('file_name', 'document.docx')}",
                caption="âœ… Semua perubahan telah diterapkan. Berikut file finalnya."
            )

            # 3. Hapus pesan loading setelah file terkirim
            await status_msg.delete()

        except Exception as e:
            # Jika error, beri tahu pengguna alih-alih menghapus pesan
            await query.message.reply_text(f"âŒ Terjadi kesalahan saat mengirim file: {e}")

        # Bersihkan data user
        context.user_data.clear()
        return ConversationHandler.END

# Definisikan ConversationHandler di sini
edit_word_v2_conv = ConversationHandler(
    entry_points=[CommandHandler("edit", mulai_edit_v2)],
    states={
        MENUNGGU_KATA_LAMA: [MessageHandler(filters.TEXT & ~filters.COMMAND, proses_cari_kata)],
        KONFIRMASI_AKSI: [CallbackQueryHandler(minta_kata_baru, pattern="^(ganti_semua|batal)$")],
        MENUNGGU_KATA_BARU: [MessageHandler(filters.TEXT & ~filters.COMMAND, eksekusi_simpan_perubahan)],
        PILIHAN_LANJUT: [CallbackQueryHandler(handle_pilihan_akhir, pattern="^(lanjut_edit|cukup_kirim)$")],
    },
    fallbacks=[CommandHandler("cancel", lambda u, c: ConversationHandler.END)],
)

